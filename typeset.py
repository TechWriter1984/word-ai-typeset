"""步骤2：LLM 分类 + 自动排版。

输入：用户前处理后的 docx 路径
行为：调用大模型对每段做样式分类，按 config.style_map 映射到
      Word 自定义样式名并应用。最后可选执行清理标题编号。
      （format_tables / format_images 统一移到 post_process 阶段，
        因为 merge 之后才能知道哪些块属于 BodyText 节，避免误处理
        封面 Logo / SymbolAndTerm 表 Icon。）
输出：排版后的 docx（保存到 cfg.output_dir，文件名加后缀 _typeset）
"""
from __future__ import annotations

import os
from typing import List, Tuple

import copy

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from config import Config
from llm import classify_paragraphs
from typeset_extra import clean_heading_numbers


def _import_styles_from_template(src_doc: Document, tpl_path: str, needed_names) -> int:
    """从模板把 needed_names 里的样式定义复制到源 docx（已存在的不覆盖）。"""
    if not os.path.exists(tpl_path):
        return 0
    tpl_doc = Document(tpl_path)
    src_styles_el = src_doc.part.styles.element
    tpl_styles_el = tpl_doc.part.styles.element

    existing = set()
    for st in src_styles_el.findall(qn("w:style")):
        nm = st.find(qn("w:name"))
        if nm is not None:
            existing.add(nm.get(qn("w:val")))

    imported = 0
    for st in tpl_styles_el.findall(qn("w:style")):
        nm = st.find(qn("w:name"))
        if nm is None:
            continue
        name = nm.get(qn("w:val"))
        if name not in needed_names or name in existing:
            continue
        src_styles_el.append(copy.deepcopy(st))
        imported += 1
    return imported


def _collect_paragraphs(doc: Document, skip_blank: bool, max_len: int) -> Tuple[List[str], List[int]]:
    """收集需要分类的段落文本及其在 doc.paragraphs 中的下标。"""
    items: List[str] = []
    indices: List[int] = []
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if skip_blank and not text:
            continue
        items.append(text[:max_len])
        indices.append(idx)
    return items, indices


def typeset_doc(cfg: Config, input_path: str) -> str:
    """对 input_path 进行 LLM 分类并应用样式，返回输出路径。"""
    os.makedirs(cfg.output_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(input_path))[0]
    out_path = os.path.join(cfg.output_dir, f"{base}_typeset.docx")

    doc = Document(input_path)

    needed = set(cfg.style_map.values())
    imported = _import_styles_from_template(doc, cfg.template_docx, needed)
    if imported:
        print(f"[typeset] 从模板导入 {imported} 个样式定义")

    items, indices = _collect_paragraphs(
        doc,
        skip_blank=cfg.typeset.get("skip_blank", True),
        max_len=cfg.typeset.get("max_text_len", 500),
    )
    print(f"[typeset] 共 {len(items)} 段待分类")

    labels = classify_paragraphs(cfg, items)
    style_map = cfg.style_map
    missing = set()
    applied = 0

    style_cache = {}
    for name in set(style_map.values()):
        try:
            style_cache[name] = doc.styles[name]
        except KeyError:
            style_cache[name] = None

    for label, idx in zip(labels, indices):
        target = style_map.get(label)
        if not target:
            missing.add(label)
            continue
        sty = style_cache.get(target)
        if sty is None:
            missing.add(target)
            continue
        para: Paragraph = doc.paragraphs[idx]
        try:
            para.style = sty
            applied += 1
        except Exception as e:
            missing.add(f"{target}({type(e).__name__})")

    # ---- typeset 阶段增强：默认只做「清理标题编号」----
    # 表格/图片排版留到 merge 后 post_process 执行（那时才能按 BodyText 节范围过滤）
    features = cfg.raw.get("features", {})
    ts_cfg = cfg.raw.get("typeset", {})

    if features.get("clean_heading_numbers", True):
        c_stats = clean_heading_numbers(doc)
        print(f"[typeset] 清理标题编号: 清理{c_stats['cleaned']} 跳过{c_stats['skipped']} 非标题{c_stats['not_heading']}")

    if ts_cfg.get("stage2_format_tables", False):
        from typeset_extra import format_tables
        t_stats = format_tables(doc, force_black_text=False)
        print(f"[typeset] (stage2) 表格排版: 总{t_stats['total']} 处理{t_stats['processed']}")
    if ts_cfg.get("stage2_format_images", False):
        from typeset_extra import format_images
        is_cn = features.get("caption_chinese", True)
        i_stats = format_images(doc, is_chinese=is_cn)
        print(f"[typeset] (stage2) 图片排版: 总{i_stats['total']} 处理{i_stats['processed']}")

    doc.save(out_path)
    print(f"[typeset] 完成 -> {out_path}  (应用 {applied}/{len(labels)} 段)")
    if missing:
        print(f"[typeset] 注意：以下标签/样式未命中：{sorted(missing)}")
        avail = []
        for s in doc.part.styles.element.findall(qn("w:style")):
            nm = s.find(qn("w:name"))
            if nm is not None:
                v = nm.get(qn("w:val")) or ""
                if v.startswith("_") and v.endswith("_hs") or v.startswith("heading"):
                    avail.append(v)
        print(f"[typeset] 源 docx 可用目标样式（前 30）: {sorted(set(avail))[:30]}")
    return out_path
